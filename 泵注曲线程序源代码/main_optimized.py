import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox, filedialog
from PIL import Image, ImageTk
import os
import threading
import glob
import time
import sys
from pickle import TRUE
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# ++++++++++++++++++++ UTILITY FUNCTIONS (for PyInstaller) ++++++++++++++++++++

def resource_path(relative_path):
    """ 获取资源的绝对路径，适用于开发环境和PyInstaller打包环境 """
    try:
        # PyInstaller 创建一个临时文件夹，并将路径存储在 _MEIPASS 中
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

# ++++++++++++++++++++ MODEL FUNCTIONS (Optimized with Pandas) ++++++++++++++++++++

def initialize_plot_settings():
    """Initializes matplotlib plot settings."""
    plt.rcParams['font.sans-serif'] = ['SimSun']
    plt.rcParams['axes.unicode_minus'] = False
    plt.rcParams['figure.figsize'] = (8, 6)
    plt.rcParams['font.size'] = 14

def find_header_row(df):
    """Finds the row index that contains the headers by looking for multiple keywords."""
    for i, row in df.iterrows():
        row_str = ' '.join(map(str, row.tolist()))
        if ("压力" in row_str or "施工压力" in row_str) and "排量" in row_str:
            return i
    return 0

def get_column_mappings(df, header_row_index):
    """Gets the mapping of data types to column names."""
    headers = df.iloc[header_row_index].astype(str)
    col_map = {}
    # Check for sand ratio vs concentration more robustly
    use_sand_ratio = any(h.strip().startswith("砂比") for h in headers)

    for col, header in headers.items():
        header_strip = header.strip()
        if header_strip.startswith("压力") or header_strip.startswith("施工压力"):
            col_map['压力'] = col
        elif header_strip.startswith("排量"):
            col_map['排量'] = col
        elif header_strip.startswith("累计液量") or header_strip.startswith("总液量"):
            col_map['累计液量'] = col
        elif use_sand_ratio and header_strip.startswith("砂比"):
            col_map['砂'] = col
        elif not use_sand_ratio and header_strip.startswith("砂浓度"):
            col_map['砂'] = col

    col_map['砂浓度转化'] = not use_sand_ratio
    return col_map

def process_data_optimized(source_path, target_path):
    """
    Processes the source excel file using pandas for high performance,
    collects valid sheets, and then writes them to the target file,
    replicating the original logic and output format.
    """
    print("数据整理开始 (Pandas优化版 - 最终修正)")
    
    try:
        xls = pd.ExcelFile(source_path)
    except FileNotFoundError:
        print(f"错误: 源文件 '{source_path}' 未找到。")
        return

    processed_sheets = {}
    for sheet_name in xls.sheet_names:
        print(f"正在处理段: {sheet_name}")
        
        df_raw = pd.read_excel(xls, sheet_name=sheet_name, header=None)
        header_row = find_header_row(df_raw)
        col_map = get_column_mappings(df_raw, header_row)

        if not all(k in col_map for k in ['压力', '排量', '累计液量', '砂']):
            print(f"警告: 在段 '{sheet_name}' 中缺少关键数据列，跳过。")
            continue

        df = df_raw.iloc[header_row + 1:].copy()
        df_processed = pd.DataFrame({
            '压力': pd.to_numeric(df[col_map['压力']], errors='coerce'),
            '排量': pd.to_numeric(df[col_map['排量']], errors='coerce'),
            '累计液量': pd.to_numeric(df[col_map['累计液量']], errors='coerce'),
            '砂': pd.to_numeric(df[col_map['砂']], errors='coerce')
        })

        df_processed.dropna(subset=['压力', '排量', '累计液量'], inplace=True)
        df_processed.reset_index(drop=True, inplace=True)

        if df_processed.empty:
            print(f"警告: 在段 '{sheet_name}' 中没有有效数据，跳过。")
            continue

        if col_map.get('砂浓度转化', False):
            df_processed['砂比'] = df_processed['砂'] / 15
        else:
            df_processed['砂比'] = df_processed['砂']
        
        # Find the row index of the last valid cumulative liquid value
        last_liquid_series = df_processed['累计液量'].dropna()
        if last_liquid_series.empty:
            print(f"警告: 在段 '{sheet_name}' 中没有有效的累计液量数据，跳过。")
            continue
        
        last_liquid_value = last_liquid_series.iloc[-1]
        # Find the first occurrence of this last value
        pump_time_row_index = df_processed[df_processed['累计液量'] == last_liquid_value].index[0]

        pump_time = pump_time_row_index
        total_liquid = last_liquid_value
        avg_flow_rate = total_liquid / pump_time if pump_time > 0 else 0
        
        shutdown_pressure_row_index = pump_time + 3
        shutdown_pressure = df_processed.loc[shutdown_pressure_row_index, '压力'] if shutdown_pressure_row_index < len(df_processed) else None

        p_i_row = pump_time + 900
        p_j_row = pump_time + 1200
        p_i = df_processed.loc[p_i_row, '压力'] if p_i_row < len(df_processed) else None
        p_j = df_processed.loc[p_j_row, '压力'] if p_j_row < len(df_processed) else None

        g_function_params = pd.DataFrame({
            '平均排量': [avg_flow_rate], '泵注时间': [pump_time], '停泵压力': [shutdown_pressure],
            'p_i': [p_i], 't_i': [p_i_row], 'p_j': [p_j], 't_j': [p_j_row]
        })

        # Replicate original output format by concatenating horizontally
        final_df = pd.concat([df_processed[['压力', '排量', '砂比', '累计液量']], g_function_params], axis=1)
        processed_sheets[sheet_name] = final_df

    if not processed_sheets:
        print("错误: 没有生成任何有效的工作表，目标文件未保存。请检查源文件格式和内容。")
        raise ValueError("没有要处理的有效数据。")
    else:
        with pd.ExcelWriter(target_path, engine='openpyxl') as writer:
            for sheet_name, df_to_write in processed_sheets.items():
                df_to_write.to_excel(writer, sheet_name=sheet_name, index=False)
        print('数据整理处理完成！')


def generate_figures_optimized(target_path):
    """Generates plots from the processed data using pandas."""
    initialize_plot_settings()
    print("图片生成开始运行 (Pandas优化版)")
    figs_path = "figs"
    os.makedirs(figs_path, exist_ok=True)

    try:
        xls = pd.ExcelFile(target_path)
    except FileNotFoundError:
        print(f"错误: 处理后的文件 '{target_path}' 未找到。")
        return

    for i, sheet_name in enumerate(xls.sheet_names):
        print(f"正在进行第{i+1}张图片处理: {sheet_name}")
        df = pd.read_excel(xls, sheet_name=sheet_name)
        
        df.dropna(subset=['压力', '排量', '砂比'], inplace=True)
        if df.empty:
            print(f"警告: 段 '{sheet_name}' 数据为空，无法生成图片。")
            continue

        time_min = df.index / 60

        fig, ax1 = plt.subplots()
        ax2 = ax1.twinx()

        ax1.plot(time_min, df['压力'], 'r', label="施工压力(MPa)")
        ax1.plot(time_min, df['排量'], 'g', label="排量(m$^3$/min)")
        ax1.plot(time_min, df['砂比'], 'black', alpha=0.5, label="砂比(%)")

        ax1.set_xlabel("泵注时间(min)")
        ax1.set_ylabel("施工压力(MPa)\n排量(m$^3$/min)")
        ax2.set_ylabel("砂比(%)")

        ax1.set_yticks(np.arange(0, 81, 10))
        ax2.set_yticks(np.arange(0, 81, 10))
        
        ax1.legend()
        ax1.grid(True)
        
        plt.savefig(os.path.join(figs_path, f'施工曲线第{i+1}段.png'), dpi=600)
        plt.close(fig)

    print("图片处理完成")

def generate_doc_optimized(target_path):
    """Generates a Word document with the created figures."""
    print("开始生成曲线文档 (Pandas优化版)")
    try:
        xls = pd.ExcelFile(target_path)
    except FileNotFoundError:
        print(f"错误: 处理后的文件 '{target_path}' 未找到。")
        return
        
    figs_path = "figs"
    doc = Document()
    
    heading = doc.add_heading(level=2)
    run = heading.add_run("3.5 施工曲线图")
    font = run.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    for i, sheet_name in enumerate(xls.sheet_names):
        fig_path = os.path.join(figs_path, f'施工曲线第{i+1}段.png')
        if not os.path.exists(fig_path):
            print(f"警告: 图片 '{fig_path}' 未找到，跳过。")
            continue
            
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        run.add_picture(fig_path, width=Cm(12))

        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(f"图3-5-{i+1}    第{i+1}级泵注曲线")
        font = run.font
        font.name = 'Times New Roman'
        font.size = Pt(10.5)
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    doc.save(os.path.join(figs_path, '施工曲线图.docx'))
    print("文档生成完成")

def process_pump_injection_curve_optimized():
    source = '泵注曲线待处理.xlsx'
    target = '泵注曲线.xlsx'
    print("\n开始处理泵注曲线数据...")
    process_data_optimized(source, target)
    print("\n开始生成泵注曲线图片...")
    generate_figures_optimized(target)
    print("\n开始生成曲线文档...")
    generate_doc_optimized(target)
    print("\n泵注曲线处理完成！")

# NOTE: The following functions for pressure-flow ratio and shutdown curve
# are kept similar to the original for now, but could also be optimized.
def YP_solve(source,target):
    wb = pd.ExcelFile(source)
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        for i, sheet_name in enumerate(wb.sheet_names):
            print(f'正在处理第{i+1}段: {sheet_name}')
            df = pd.read_excel(wb, sheet_name=sheet_name)
            df_filtered = df[pd.to_numeric(df['砂比'], errors='coerce') > 0]
            df_filtered.to_excel(writer, sheet_name=str(i+1), index=False)
    print('压排比数据处理完成！')
    
def YP_fig(target):
    initialize_plot_settings()
    print("压排比图片生成开始运行")
    figs_path = "figs"
    os.makedirs(figs_path, exist_ok=True)

    xls = pd.ExcelFile(target)
    for i, sheet_name in enumerate(xls.sheet_names):
        print(f"正在进行第{i+1}张压排比图片处理")
        df = pd.read_excel(xls, sheet_name=sheet_name)
        
        df['压力'] = pd.to_numeric(df['压力'], errors='coerce')
        df['排量'] = pd.to_numeric(df['排量'], errors='coerce')
        df['砂比'] = pd.to_numeric(df['砂比'], errors='coerce')
        df.dropna(subset=['压力', '排量', '砂比'], inplace=True)

        min_val = df.index / 60
        # Avoid division by zero
        pressure_flow_ratio = df['压力'].divide(df['排量']).replace([np.inf, -np.inf], 0)

        fig, ax1 = plt.subplots()
        ax2 = ax1.twinx()
        ax1.plot(min_val, pressure_flow_ratio, 'r', label="压排比(MPa/(m$^3$/min))")
        ax2.plot(min_val, df['砂比'], 'black', alpha=0.5, label="砂比(%)")
        ax1.set_xlabel("泵注时间(min)")
        ax1.set_ylabel("压力(MPa)/排量(m$^3$/min)")
        ax2.set_ylabel("砂比(%)")
        ax1.set_ylim(0, 5)
        ax2.set_ylim(0, 100)
        ax1.set_yticks(np.arange(0, 5.1, 0.5))
        ax2.set_yticks(np.arange(0, 101, 10))
        
        lines, labels = ax1.get_legend_handles_labels()
        lines2, labels2 = ax2.get_legend_handles_labels()
        ax2.legend(lines + lines2, labels + labels2, loc=0)

        ax1.grid(True)
        plt.savefig(os.path.join(figs_path, f'压排比曲线第{i+1}段.png'), dpi=600)
        plt.close(fig)
    print("压排比图片处理完成")

def TBsolve(source):
    target = "停泵曲线.xlsx"
    try:
        source_xls = pd.ExcelFile(source)
    except FileNotFoundError:
        print(f"文件 {source} 不存在")
        return

    all_shutdown_data = []
    for j, sheet_name in enumerate(source_xls.sheet_names):
        print(f"处理第 {j+1} 段停泵数据")
        df = pd.read_excel(source_xls, sheet_name=sheet_name)
        
        # Correctly get t_i from the first row where it's not NaN
        start_row_series = df['t_i'].dropna()
        if not start_row_series.empty:
            start_row_val = int(start_row_series.iloc[0])
            
            # The pressure data for shutdown curve is in the main '压力' column
            # of the processed sheet, starting from t_i.
            if start_row_val < len(df):
                # Extract the pressure data from the '压力' column (already cleaned)
                shutdown_pressures = df['压力'].iloc[start_row_val:].tolist()
                all_shutdown_data.extend(shutdown_pressures)
                all_shutdown_data.extend([None] * 50)  # Spacer
            else:
                print(f"第 {j+1} 段的停泵起始行号 {start_row_val} 超出范围，跳过。")
        else:
            print(f"第 {j+1} 段未找到有效的停泵起始行号 (t_i)，跳过。")

    if not all_shutdown_data:
        print("警告: 未能从任何段中提取到停泵数据。")
        df_shutdown = pd.DataFrame({'时间': [], '压力': []})
    else:
        df_shutdown = pd.DataFrame({'压力': all_shutdown_data})
        df_shutdown['时间'] = range(1, len(df_shutdown) + 1)
        df_shutdown = df_shutdown[['时间', '压力']]
        df_shutdown['压力'] = df_shutdown['压力'].replace(0, np.nan)

    df_shutdown.to_excel(target, index=False)
    print('停泵数据处理完成')

def TBfig(source):
    initialize_plot_settings()
    try:
        df = pd.read_excel(source)
    except FileNotFoundError:
        print(f"文件 {source} 不存在")
        return
        
    # --- Logic to remove trailing data points from each segment ---
    processed_data = []
    segment_start_index = 0
    # Find indices where a segment ends (NaN value)
    nan_indices = df[df['压力'].isna()].index.tolist()
    
    # Add the end of the dataframe as the final boundary
    if not nan_indices or nan_indices[-1] < len(df) - 1:
        nan_indices.append(len(df))

    for end_index in nan_indices:
        segment = df.iloc[segment_start_index:end_index]
        if len(segment) > 10:
            # Remove last 10 points
            processed_data.append(segment.iloc[:-10])
        else:
            # Keep short segments as is
            processed_data.append(segment)
        segment_start_index = end_index + 1
    
    if not processed_data:
        print("警告: 处理后没有有效的停泵数据用于绘图。")
        return

    fig, ax1 = plt.subplots()
    
    # Plot each segment individually to create visual separation
    for i, segment in enumerate(processed_data):
        segment_cleaned = segment.dropna()
        # Add label only to the first segment to avoid duplicate legend entries
        label = "压力(MPa)" if i == 0 else ""
        ax1.plot(segment_cleaned['时间'], segment_cleaned['压力'], 'r', label=label)

    ax1.set_xlabel("时间(s)")
    ax1.set_ylabel("压力(MPa)")
    ax1.set_yticks(np.arange(0, 51, 5))
    ax1.set_title("压降数据（MPa）")
    ax1.legend()
    ax1.grid(True)
    plt.savefig('figs/停泵曲线.png', dpi=600)
    print("停泵曲线图片处理完成")

def process_pressure_flow_ratio_optimized():
    source = '泵注曲线.xlsx'
    target = '压排比.xlsx'
    print("\n开始处理压排比数据...")
    YP_solve(source, target)
    print("\n开始生成压排比图片...")
    YP_fig(target)
    print("\n压排比处理完成！")

def process_shutdown_curve_optimized():
    source = '泵注曲线.xlsx'
    target = '停泵曲线.xlsx'
    print("\n开始处理停泵曲线数据...")
    TBsolve(source)
    print("\n开始生成停泵曲线图片...")
    TBfig(target)
    print("\n停泵曲线处理完成！")

# ++++++++++++++++++++ GUI APPLICATION CLASS ++++++++++++++++++++

class Application(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("压裂数据处理系统 (Pandas优化版)")
        self.geometry("1200x700")
        self.source_file_path = None

        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=2)

        self.control_frame = ttk.Frame(self, padding="10")
        self.control_frame.grid(row=0, column=0, sticky="nsew")
        self.control_frame.grid_rowconfigure(6, weight=1)

        self.btn_select_file = ttk.Button(self.control_frame, text="选择源文件", command=self.select_file)
        self.btn_select_file.grid(row=0, column=0, pady=5, sticky="ew")

        self.file_path_label = ttk.Label(self.control_frame, text="未选择文件", anchor="w", relief="sunken")
        self.file_path_label.grid(row=1, column=0, pady=2, sticky="ew")

        self.btn_pump_injection = ttk.Button(self.control_frame, text="处理泵注曲线", command=lambda: self.run_function_in_thread(self.process_pump_injection_curve_gui))
        self.btn_pump_injection.grid(row=2, column=0, pady=5, sticky="ew")

        self.btn_pressure_flow = ttk.Button(self.control_frame, text="处理压排比", command=lambda: self.run_function_in_thread(self.process_pressure_flow_ratio_gui))
        self.btn_pressure_flow.grid(row=3, column=0, pady=5, sticky="ew")

        self.btn_shutdown_curve = ttk.Button(self.control_frame, text="处理停泵曲线", command=lambda: self.run_function_in_thread(self.process_shutdown_curve_gui))
        self.btn_shutdown_curve.grid(row=4, column=0, pady=5, sticky="ew")

        self.btn_full_process = ttk.Button(self.control_frame, text="全流程处理", command=lambda: self.run_function_in_thread(self.process_full_gui))
        self.btn_full_process.grid(row=5, column=0, pady=5, sticky="ew")

        self.btn_exit = ttk.Button(self.control_frame, text="退出程序", command=self.quit_app)
        self.btn_exit.grid(row=6, column=0, pady=5, sticky="ew")

        self.log_text = scrolledtext.ScrolledText(self.control_frame, wrap=tk.WORD, height=15)
        self.log_text.grid(row=7, column=0, pady=10, sticky="nsew")
        self.log_text.config(state=tk.DISABLED)

        sys.stdout = ConsoleRedirector(self.log_text)
        sys.stderr = ConsoleRedirector(self.log_text)

        self.image_frame = ttk.Frame(self, padding="10")
        self.image_frame.grid(row=0, column=1, sticky="nsew")
        self.image_frame.grid_rowconfigure(0, weight=1)
        self.image_frame.grid_columnconfigure(0, weight=1)

        self.image_canvas = tk.Canvas(self.image_frame, bg="lightgray")
        self.image_canvas.grid(row=0, column=0, sticky="nsew")

        self.image_files = []
        self.image_combobox = ttk.Combobox(self.image_frame, values=self.image_files, state="readonly")
        self.image_combobox.grid(row=1, column=0, pady=10, sticky="ew")
        self.image_combobox.bind("<<ComboboxSelected>>", self.display_selected_image)

        self.set_buttons_state(tk.DISABLED)
        self.update_image_list()

    def select_file(self):
        file_path = filedialog.askopenfilename(
            title="选择 '泵注曲线待处理.xlsx' 文件",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if file_path:
            if os.path.basename(file_path) != '泵注曲线待处理.xlsx':
                messagebox.showwarning("文件名错误", "请选择名为 '泵注曲线待处理.xlsx' 的文件。")
                return
            self.source_file_path = file_path
            self.file_path_label.config(text=file_path)
            self.set_buttons_state(tk.NORMAL)
            self.log_message(f"已选择文件: {file_path}")
            self.update_image_list()

    def run_function_in_thread(self, func):
        if not self.source_file_path:
            messagebox.showerror("错误", "请先选择一个源文件！")
            return

        self.log_text.config(state=tk.NORMAL)
        self.log_text.delete(1.0, tk.END)
        self.log_text.config(state=tk.DISABLED)
        self.set_buttons_state(tk.DISABLED)
        self.log_message(f"开始执行...")
        
        thread = threading.Thread(target=self.run_with_changed_dir, args=(func,))
        thread.start()

    def run_with_changed_dir(self, func):
        original_dir = os.getcwd()
        target_dir = os.path.dirname(self.source_file_path)
        try:
            os.chdir(target_dir)
            func()
        finally:
            os.chdir(original_dir)
            self.after(0, self.set_buttons_state, tk.NORMAL)

    def set_buttons_state(self, state):
        self.btn_pump_injection.config(state=state)
        self.btn_pressure_flow.config(state=state)
        self.btn_shutdown_curve.config(state=state)
        self.btn_full_process.config(state=state)

    def log_message(self, message):
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)

    def update_image_list(self):
        if not self.source_file_path:
            self.image_files = []
        else:
            figs_dir = os.path.join(os.path.dirname(self.source_file_path), "figs")
            if not os.path.exists(figs_dir):
                os.makedirs(figs_dir)
            image_paths = sorted(glob.glob(os.path.join(figs_dir, "*.png")))
            self.image_files = [os.path.basename(f) for f in image_paths]

        self.image_combobox['values'] = self.image_files
        if self.image_files:
            self.image_combobox.set(self.image_files[-1])
            self.display_selected_image()
        else:
            self.image_combobox.set("")
            self.image_canvas.delete("all")

    def display_selected_image(self, event=None):
        selected_image_name = self.image_combobox.get()
        if selected_image_name and self.source_file_path:
            image_path = os.path.join(os.path.dirname(self.source_file_path), "figs", selected_image_name)
            self.display_image(image_path)

    def display_image(self, image_path):
        try:
            self.image_canvas.update_idletasks()
            canvas_width = self.image_canvas.winfo_width()
            canvas_height = self.image_canvas.winfo_height()

            if canvas_width <= 1 or canvas_height <= 1:
                self.after(100, lambda: self.display_image(image_path))
                return

            img = Image.open(image_path)
            img_width, img_height = img.size
            
            ratio = min(canvas_width / img_width, canvas_height / img_height)
            new_width = int(img_width * ratio * 0.95)
            new_height = int(img_height * ratio * 0.95)

            img = img.resize((new_width, new_height), Image.LANCZOS)
            self.photo = ImageTk.PhotoImage(img)
            
            self.image_canvas.delete("all")
            x = (canvas_width - new_width) / 2
            y = (canvas_height - new_height) / 2
            self.image_canvas.create_image(x, y, anchor=tk.NW, image=self.photo)
        except Exception as e:
            self.log_message(f"显示图片时发生错误：{e}")
            self.image_canvas.delete("all")

    def process_pump_injection_curve_gui(self):
        try:
            process_pump_injection_curve_optimized()
            self.log_message("泵注曲线处理完成！")
        except Exception as e:
            self.log_message(f"泵注曲线处理失败：{e}")
            messagebox.showerror("错误", f"泵注曲线处理失败：{e}")
        finally:
            self.after(0, self.update_image_list)

    def process_pressure_flow_ratio_gui(self):
        try:
            process_pressure_flow_ratio_optimized()
            self.log_message("压排比处理完成！")
        except Exception as e:
            self.log_message(f"压排比处理失败：{e}")
            messagebox.showerror("错误", f"压排比处理失败：{e}")
        finally:
            self.after(0, self.update_image_list)

    def process_shutdown_curve_gui(self):
        try:
            process_shutdown_curve_optimized()
            self.log_message("停泵曲线处理完成！")
        except Exception as e:
            self.log_message(f"停泵曲线处理失败：{e}")
            messagebox.showerror("错误", f"停泵曲线处理失败：{e}")
        finally:
            self.after(0, self.update_image_list)

    def process_full_gui(self):
        try:
            self.log_message("开始全流程处理...")
            process_pump_injection_curve_optimized()
            process_pressure_flow_ratio_optimized()
            process_shutdown_curve_optimized()
            self.log_message("全流程处理完成！")
        except Exception as e:
            self.log_message(f"全流程处理失败：{e}")
            messagebox.showerror("错误", f"全流程处理失败：{e}")
        finally:
            self.after(0, self.update_image_list)

    def quit_app(self):
        self.log_message("程序已退出。")
        self.destroy()
        sys.exit(0)

class ConsoleRedirector:
    def __init__(self, text_widget):
        self.text_widget = text_widget
        self.stdout = sys.__stdout__
        self.stderr = sys.__stderr__

    def write(self, message):
        self.text_widget.after(0, self.insert_text, message)
        if self.stdout:
            self.stdout.write(message)
            self.stdout.flush()

    def insert_text(self, message):
        self.text_widget.config(state=tk.NORMAL)
        self.text_widget.insert(tk.END, message)
        self.text_widget.see(tk.END)
        self.text_widget.config(state=tk.DISABLED)

    def flush(self):
        if self.stdout:
            self.stdout.flush()

if __name__ == "__main__":
    app = Application()
    app.mainloop()
